"""
Build the Unit 19 Business Intelligence assignment report (Task 1).
Output: AdventureWorks_BI_Assignment_Report.docx
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- Palette ----------
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x5E, 0x8C)
GOLD = RGBColor(0xC8, 0xA0, 0x4B)
GREY = RGBColor(0x55, 0x55, 0x55)
DARK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()

# ---------- Base styles ----------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=None, size=10, align="left", italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return p

def heading(text, level=1):
    if level == 1:
        p = doc.add_paragraph()
        p.space_before = Pt(14)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = NAVY
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        # bottom border
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "C8A04B")
        pbdr.append(bottom)
        pPr.append(pbdr)
        return p
    elif level == 2:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = BLUE
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        return p
    else:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.color.rgb = DARK
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        return p

def body(text, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    p.paragraph_format.space_after = Pt(space_after)
    return p

def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(4)
    return p

# =====================================================================
# TITLE PAGE
# =====================================================================
for _ in range(4):
    doc.add_paragraph()

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Unit 19: Business Intelligence")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Business Intelligence Dashboard for Adventure Works Cycles")
r.font.size = Pt(15)
r.font.color.rgb = BLUE

doc.add_paragraph()
band = doc.add_paragraph()
band.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = band.add_run("TASK 1 — Business Context and BI Strategy")
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = GOLD

for _ in range(6):
    doc.add_paragraph()

meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ("Qualification", "BTEC Higher National — Computing"),
    ("Unit", "Unit 19: Business Intelligence"),
    ("Assignment", "BI Dashboard for Adventure Works Cycles"),
    ("Learning Aims Covered", "LO1, LO2 (with foundations for LO3 & LO4)"),
]
for i, (k, v) in enumerate(meta_data):
    set_cell_text(meta.rows[i].cells[0], k, bold=True, color=NAVY, size=11)
    set_cell_text(meta.rows[i].cells[1], v, size=11)
meta.columns[0].width = Inches(2.4)
meta.columns[1].width = Inches(3.6)

doc.add_page_break()

# =====================================================================
# CONTENTS
# =====================================================================
heading("Contents", 1)
toc_items = [
    "1.  What is Business Intelligence and its Role in Retail",
    "2.  Adventure Works Business Processes and Data Flow",
    "3.  The Five Most Important Business Questions and KPIs",
    "4.  BI Tool Recommendation: Power BI vs Tableau vs Qlik Sense",
    "5.  Data Categories: Structured, Semi-Structured and Unstructured",
    "References",
]
for item in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(item)
    r.font.size = Pt(11.5)

doc.add_page_break()

# =====================================================================
# SECTION 1
# =====================================================================
heading("1.  What is Business Intelligence and its Role in Retail", 1)

body(
    "Business Intelligence (BI) is the combination of strategies, technologies, processes and "
    "architectures that an organisation uses to collect raw data from its operational systems, "
    "transform it into clean and consistent information, and present it in a form that supports "
    "fact-based decision making. In practical terms, BI turns the transactions a business records "
    "every day — sales orders, returns, stock movements, customer registrations — into the "
    "dashboards, reports and key performance indicators (KPIs) that managers actually use to run "
    "the business. Gartner describes BI as an umbrella term covering the applications, "
    "infrastructure, tools and best practices that enable access to and analysis of information to "
    "improve and optimise decisions and performance (Gartner, 2023)."
)
body(
    "A widely used way of understanding BI is to separate it into three decision levels, because "
    "different stakeholders consume information at different speeds and depths:"
)
bullet(" — day-to-day questions answered with near real-time data, e.g. \"how many orders were placed yesterday?\" or \"which product line is currently selling fastest?\".", bold_lead="Operational BI")
bullet(" — month-on-month and quarter-on-quarter analysis used by middle managers to adjust tactics, e.g. comparing regional performance or monitoring return rates against a target.", bold_lead="Tactical BI")
bullet(" — long-range trend analysis used by executives to set direction, e.g. multi-year revenue growth, profitability of whole categories, or customer-segment value.", bold_lead="Strategic BI")

heading("The role of BI in retail today", 2)
body(
    "Retail is one of the most data-intensive sectors in the economy. Every checkout, every product "
    "page view, every return and every loyalty sign-up generates data, and the margins in retail are "
    "typically thin, so the difference between profit and loss is often found in small operational "
    "efficiencies. BI plays several critical roles in a modern retailer:"
)
bullet(" — a single, governed version of the truth replaces conflicting spreadsheet exports, so executives and managers argue about decisions rather than about whose numbers are correct.", bold_lead="A single source of truth")
bullet(" — analysing sales by product, territory and time allows retailers to stock the right goods in the right place, reducing both stock-outs and costly over-stock.", bold_lead="Demand and inventory optimisation")
bullet(" — demographic and behavioural analysis lets the business identify its most valuable customers and tailor offers, improving retention and lifetime value.", bold_lead="Customer insight and segmentation")
bullet(" — return rates, profit margins and average order value are monitored continuously so problems (such as a faulty product line generating excessive returns) are caught early.", bold_lead="Margin and returns control")

heading("Adventure Works as the reference point", 2)
body(
    "Adventure Works Cycles is a multinational bicycle manufacturer and retailer operating across "
    "Canada, the United Kingdom, France, Australia and the United States, selling road, mountain "
    "and touring bikes plus accessories through its own four stores and partner retailers. Three "
    "years after operations began, the Chief Commercial Officer (CCO) is no longer willing to base "
    "decisions on manual spreadsheet exports and has commissioned a Power BI solution to act as the "
    "company's single version of the truth. The 2020–2022 dataset that Adventure Works provides is a "
    "classic retail data estate: a Sales fact table, a Returns fact table, and dimension tables "
    "describing customers, products (with categories and subcategories), territories and a full "
    "calendar. Throughout this report Adventure Works is used as the concrete example: the BI "
    "strategy described here is the one that will drive the four-page dashboard built in Task 3, and "
    "the KPIs defined in Section 3 are exactly the measures the executives and managers have asked for."
)

# =====================================================================
# SECTION 2
# =====================================================================
heading("2.  Adventure Works Business Processes and Data Flow", 1)
body(
    "Before any dashboard can be designed it is essential to understand which business processes "
    "generate the data and how that data flows through the organisation into the BI layer. A "
    "business process is a repeatable sequence of activities that creates value; in Adventure Works "
    "the core processes that feed the data warehouse are described below."
)

heading("Core business processes", 2)
processes = [
    ("Customer Acquisition", "New customers register or are on-boarded by partner retailers. This process creates and maintains the Customer Lookup dimension, including gender, occupation, education level, marital status and annual income."),
    ("Order Management", "When a customer places an order, the order header and line items are captured. This is the heart of the operation and writes rows into the Sales Data fact table (OrderDate, CustomerKey, ProductKey, TerritoryKey, OrderQuantity, TotalRevenue)."),
    ("Sales and Distribution", "Confirmed orders are priced, fulfilled and dispatched through Adventure Works stores or partner retailers across the five countries. This process enriches the Sales fact and links it to the Territory dimension."),
    ("Inventory and Fulfilment", "Stock is manufactured, held and picked against orders. Product cost and price held in the Product Lookup dimension support profitability analysis and feed inventory decisions."),
    ("Returns Management", "Where goods are sent back, the return is logged (ReturnDate, ProductKey, TerritoryKey, ReturnQuantity), populating the Returns Data fact table and enabling return-rate and quality analysis."),
    ("Customer Experience", "Post-sale interactions and the resulting demographic and behavioural profile support segmentation and retention analysis for the Customer Insights team."),
]
ptab = doc.add_table(rows=1, cols=2)
ptab.style = "Light Grid Accent 1"
hdr = ptab.rows[0].cells
set_cell_text(hdr[0], "Business Process", bold=True, color=WHITE, size=10.5)
set_cell_text(hdr[1], "What it does and which data it generates", bold=True, color=WHITE, size=10.5)
shade_cell(hdr[0], "1F3A5F")
shade_cell(hdr[1], "1F3A5F")
for name, desc in processes:
    row = ptab.add_row().cells
    set_cell_text(row[0], name, bold=True, color=NAVY, size=10)
    set_cell_text(row[1], desc, size=10)
ptab.columns[0].width = Inches(1.8)
ptab.columns[1].width = Inches(4.4)

heading("How data is pushed through each process", 2)
body(
    "Each operational process is a data source. Source data is extracted and loaded into Power Query, "
    "where it is cleaned and transformed (the ETL stage), and then organised into a star schema data "
    "warehouse model with two fact tables (Sales and Returns) surrounded by conformed dimension tables "
    "(Customer, Product, Territory, Calendar). The Power BI dashboard then sits on top of this model and "
    "serves the four stakeholder groups. The diagram below shows this end-to-end flow."
)

# Insert the diagram
doc.add_picture("diagram_process_flow.png", width=Inches(6.3))
cap = doc.paragraphs[-1]
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
fig = doc.add_paragraph()
fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fig.add_run("Figure 1 — Adventure Works business process and data-flow diagram. "
                "Operational processes feed the ETL/data-warehouse layer, which serves the Power BI dashboard.")
r.italic = True
r.font.size = Pt(9)
r.font.color.rgb = GREY

body(
    "Reading the diagram from top to bottom: the operational processes (blue) each act as a system of "
    "record and continuously generate transactions. These flow down (dashed arrows) into the ETL / data "
    "warehouse layer (gold), where Power Query cleans and conforms them into the star schema. Finally "
    "the modelled data is published to the Power BI presentation layer (green), which is divided into "
    "the four stakeholder pages. This separation of the operational layer, the warehouse layer and the "
    "presentation layer is the standard dimensional architecture described by Kimball and Ross (2013) "
    "and is what allows a single trusted model to answer many different business questions.",
    space_after=6,
)

# =====================================================================
# SECTION 3
# =====================================================================
heading("3.  The Five Most Important Business Questions and KPIs", 1)
body(
    "The dashboard is designed around the information needs of four stakeholder groups. The five KPIs "
    "below were selected because each one answers a specific, high-value business question for a named "
    "stakeholder, and together they cover revenue, profitability, growth, regional performance and "
    "product quality. Each KPI is implemented as a DAX measure (developed fully in Task 2) so that it "
    "responds correctly to slicers and cross-filtering on every page."
)

kpis = [
    ("Total Revenue",
     "Chief Commercial Officer",
     "How much money is the business generating overall, and how does it compare with previous years?",
     "SUM of TotalRevenue from the Sales fact table. Displayed as a KPI card on the Executive Overview page and trended monthly across 2020–2022.",
     "Revenue is the headline figure for the CCO; every other measure is judged in its context."),
    ("Profit Margin %",
     "Chief Commercial Officer",
     "Are we growing profitably, or are we buying revenue at the cost of margin?",
     "(Total Revenue − Total Cost) / Total Revenue, where Total Cost = SUMX(Sales, OrderQuantity × ProductCost). Shown as a card and decomposed by subcategory.",
     "Revenue alone can be misleading; margin shows whether growth is healthy and protects the thin margins typical of retail."),
    ("Revenue by Territory (Country / Continent)",
     "Regional Sales Manager",
     "Which countries and regions are over- or under-performing so I can act on the right territory?",
     "Total Revenue sliced by the Territory dimension (continent → country → region) and visualised on a filled map and a clustered bar chart.",
     "Allows the regional manager to compare like-for-like performance and target interventions geographically rather than in aggregate."),
    ("Return Rate %",
     "Product Manager",
     "Which products are being returned too often, indicating a quality or fit problem?",
     "Total Return Quantity / Total Order Quantity per product, charted against a 5% target reference line.",
     "A high return rate erodes margin and signals product issues; flagging it early protects profitability and brand reputation."),
    ("Revenue per Customer (with Top-N analysis)",
     "Customer Insights Team",
     "Who are our most valuable customers and which demographic segments drive the most revenue?",
     "Total Revenue / DISTINCTCOUNT(CustomerKey), combined with Top-20 customer ranking and revenue broken down by occupation, education, gender and marital status.",
     "Identifies high-value segments so marketing and retention spend is focused where it generates the greatest lifetime value."),
]

for i, (name, stake, q, calc, why) in enumerate(kpis, start=1):
    heading(f"KPI {i}: {name}", 3)
    tab = doc.add_table(rows=4, cols=2)
    tab.style = "Light List Accent 1"
    rows = [
        ("Stakeholder", stake),
        ("Business question answered", q),
        ("How it is calculated", calc),
        ("Why it matters", why),
    ]
    for j, (k, v) in enumerate(rows):
        set_cell_text(tab.rows[j].cells[0], k, bold=True, color=NAVY, size=10)
        set_cell_text(tab.rows[j].cells[1], v, size=10)
    tab.columns[0].width = Inches(1.9)
    tab.columns[1].width = Inches(4.3)
    doc.add_paragraph()

# =====================================================================
# SECTION 4
# =====================================================================
heading("4.  BI Tool Recommendation: Power BI vs Tableau vs Qlik Sense", 1)
body(
    "Adventure Works must select a single enterprise BI platform. The three credible candidates are "
    "Microsoft Power BI, Salesforce Tableau and Qlik Sense (Google Looker Studio is excluded as it is "
    "not positioned for enterprise-scale governance). The comparison below evaluates each tool against "
    "the factors that matter for Adventure Works specifically: the scale of a multinational retailer, a "
    "clean star-schema data structure, and a mixed audience of executives, managers and analysts."
)

comp = doc.add_table(rows=1, cols=4)
comp.style = "Light Grid Accent 1"
ch = comp.rows[0].cells
for idx, txt in enumerate(["Criterion", "Power BI", "Tableau", "Qlik Sense"]):
    set_cell_text(ch[idx], txt, bold=True, color=WHITE, size=10)
    shade_cell(ch[idx], "1F3A5F")
comp_rows = [
    ("Data modelling fit",
     "Excellent for star schemas; native relationships, DAX time-intelligence and Power Query ETL built in.",
     "Strong visuals but modelling is lighter; complex models often pushed back to the source.",
     "Strong associative engine; good for exploration but a different modelling paradigm."),
    ("Cost / scale",
     "Low entry cost (Pro per-user); Premium capacity scales to enterprise.",
     "Higher per-user licensing; can become expensive at scale.",
     "Mid-to-high licensing; capacity-based options."),
    ("Ease of use for mixed audience",
     "Familiar Office-style UI; gentle learning curve for managers, deep for analysts.",
     "Best-in-class for analyst-led visual exploration; steeper for casual users.",
     "Powerful but the associative model takes time to learn."),
    ("Ecosystem / integration",
     "Deep integration with Excel, Azure, SQL Server and Microsoft 365.",
     "Broad connectors; strong standalone analytics community.",
     "Good connectivity; smaller ecosystem than Microsoft."),
    ("Governance & single source of truth",
     "Workspaces, datasets, row-level security and the Power BI Service support a governed single source of truth.",
     "Tableau Server/Cloud provides governance; certified data sources available.",
     "Centralised governance available via Qlik Sense Enterprise."),
]
for crit, pbi, tab_, qlik in comp_rows:
    row = comp.add_row().cells
    set_cell_text(row[0], crit, bold=True, color=NAVY, size=9.5)
    set_cell_text(row[1], pbi, size=9.5)
    set_cell_text(row[2], tab_, size=9.5)
    set_cell_text(row[3], qlik, size=9.5)
comp.columns[0].width = Inches(1.3)
for c in range(1, 4):
    comp.columns[c].width = Inches(1.65)

doc.add_paragraph()
heading("Recommendation and critical evaluation", 2)
body(
    "Recommendation: Power BI is the most appropriate tool for Adventure Works. The justification is "
    "strongest when examined across the three decision levels introduced in Section 1:"
)
bullet(" — Power BI's tight integration with Excel and the Microsoft stack means store and regional staff can drill into near-real-time order and return data with almost no retraining, because the interface and formula concepts already feel familiar. Tableau and Qlik would both impose a steeper learning curve at this level.", bold_lead="At the operational level")
bullet(" — the provided data is already a clean star schema, which is exactly where Power BI's relationship model and DAX time-intelligence (year-to-date, month-on-month, year-on-year) excel. Regional and product managers get powerful tactical analysis with comparatively little engineering effort. Tableau is superb for visual exploration but leans on the underlying source for heavy modelling, and Qlik's associative engine, while powerful, is a different paradigm that the team would have to learn.", bold_lead="At the tactical level")
bullet(" — the Power BI Service provides governed workspaces, certified datasets and row-level security, giving the CCO the genuine single source of truth that was the original business driver, and Premium capacity scales economically across the five countries. Tableau can match this governance but at materially higher licensing cost, and Qlik sits in a similar cost band with a smaller ecosystem.", bold_lead="At the strategic level")
body(
    "Considering Adventure Works' scale (multinational but mid-sized), its already-dimensional data "
    "structure, and its mixed stakeholder base, Power BI offers the best balance of low total cost of "
    "ownership, the fastest route to a governed single source of truth, and the strongest native fit "
    "for star-schema modelling and time-intelligence. Industry analysis consistently places all three "
    "tools in the Leaders quadrant of Gartner's Magic Quadrant for Analytics and BI Platforms, so this "
    "is a recommendation about fit rather than about raw capability (Gartner, 2023; Howson et al., 2022). "
    "For Adventure Works specifically, Power BI is the recommended platform."
)

# =====================================================================
# SECTION 5
# =====================================================================
heading("5.  Data Categories: Structured, Semi-Structured and Unstructured", 1)
body(
    "Data is commonly classified into three categories according to how rigidly it is organised. "
    "Understanding the distinction matters because it determines how data can be stored, cleaned and "
    "queried — and the Adventure Works BI solution is built almost entirely on structured data, which "
    "is one reason a dimensional model and Power BI are such a good fit."
)

dcat = doc.add_table(rows=1, cols=3)
dcat.style = "Light Grid Accent 1"
dh = dcat.rows[0].cells
for idx, txt in enumerate(["Category", "Definition", "Examples in Adventure Works"]):
    set_cell_text(dh[idx], txt, bold=True, color=WHITE, size=10)
    shade_cell(dh[idx], "1F3A5F")
dcat_rows = [
    ("Structured",
     "Highly organised data that fits a fixed schema of rows and columns and is easily stored in relational tables and queried with SQL.",
     "The entire provided dataset: the Sales and Returns fact tables and the Customer, Product, Territory and Calendar dimension tables. Fields such as OrderDate, OrderQuantity, TotalRevenue, ProductCost and AnnualIncome are all structured."),
    ("Semi-structured",
     "Data that does not fit a strict relational schema but carries tags or markers (keys, fields) that give it some self-describing structure.",
     "A JSON or XML order feed exported from an e-commerce platform, CSV exports before they are typed in Power Query, or partner-retailer sales files whose columns vary slightly. Power Query is used to impose structure on these before loading."),
    ("Unstructured",
     "Data with no predefined model — free text, images, audio or video — which cannot be queried directly without further processing.",
     "Customer review text and free-text complaint notes, product photographs on the website, and call-centre recordings. These are not in the provided dataset but would inform future sentiment or returns-reason analysis."),
]
for cat, defn, ex in dcat_rows:
    row = dcat.add_row().cells
    set_cell_text(row[0], cat, bold=True, color=NAVY, size=10)
    set_cell_text(row[1], defn, size=9.5)
    set_cell_text(row[2], ex, size=9.5)
dcat.columns[0].width = Inches(1.2)
dcat.columns[1].width = Inches(2.4)
dcat.columns[2].width = Inches(2.6)

doc.add_paragraph()
body(
    "In summary, the Adventure Works dataset is predominantly structured, which makes it ideally suited "
    "to a relational star-schema model in Power BI. Semi-structured sources (such as raw CSV or JSON "
    "exports) are handled in the Power Query preparation stage covered in Task 2, while unstructured "
    "sources such as review text and product imagery represent a clear opportunity for future expansion "
    "of the BI capability once the core dashboard is delivered."
)

# =====================================================================
# REFERENCES
# =====================================================================
heading("References", 1)
refs = [
    "Ferrari, A. and Russo, M. (2019) The Definitive Guide to DAX. 2nd edn. Redmond: Microsoft Press.",
    "Gartner (2023) Magic Quadrant for Analytics and Business Intelligence Platforms. Stamford: Gartner, Inc.",
    "Howson, C., Richardson, J., Sallam, R. and Kronz, A. (2022) Analytics and Business Intelligence Platforms: Market Guidance. Stamford: Gartner, Inc.",
    "Jeston, J. and Nelis, J. (2014) Business Process Management. 3rd edn. Abingdon: Routledge.",
    "Kimball, R. and Ross, M. (2013) The Data Warehouse Toolkit: The Definitive Guide to Dimensional Modeling. 3rd edn. Indianapolis: Wiley.",
    "Marr, B. (2015) Big Data: Using Smart Big Data, Analytics and Metrics to Make Better Decisions and Improve Performance. Chichester: John Wiley & Sons.",
    "Microsoft (2024) Power BI Documentation. Available at: https://learn.microsoft.com/en-us/power-bi/ (Accessed: 3 June 2026).",
    "Sharda, R., Delen, D. and Turban, E. (2018) Business Intelligence, Analytics, and Data Science: A Managerial Perspective. 4th edn. Harlow: Pearson.",
]
for r in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(r)
    run.font.size = Pt(10.5)

# ---------- Footer with page numbers ----------
section = doc.sections[0]
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("Unit 19 Business Intelligence — Task 1 — Adventure Works Cycles    |    Page ")
run.font.size = Pt(8)
run.font.color.rgb = GREY
# page number field
fld1 = OxmlElement("w:fldSimple")
fld1.set(qn("w:instr"), "PAGE")
fp._p.append(fld1)

doc.save("AdventureWorks_BI_Assignment_Report.docx")
print("Saved AdventureWorks_BI_Assignment_Report.docx")
